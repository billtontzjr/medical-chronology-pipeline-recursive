"""Uncertain diagnostics are visible in review, never silently approved."""
import copy
import hashlib
import json
import logging

from docx import Document
import pytest

from src.chronology_agent import ChronologyAgent
from src.chronology_scope import ScopeReviewRequired, parse_scoped_response
from src.manual_review import parse_draft_response, review_markdown, DRAFT_LABEL
from src.scope_recovery import screen_batch


def fixture():
    docs = [
        {'filename': 'uncertain.txt', 'source_file': 'uncertain.txt', 'source_pages': [7],
         'content': '=== SOURCE PDF PAGE 7 ===\nCollected On 01/12/26 0900\n'
                    'Facility: Synthetic Imaging\nProvider: Example, MD\n'
                    'Study: Lumbar MRI\nImpression: No acute fracture.'},
        {'filename': 'supported.txt', 'content': 'Date of service: 01/19/2026\n'
         'Facility: Synthetic Imaging\nProvider: Example, MD\nStudy: Lumbar MRI\n'
         'Impression: No canal stenosis.'},
    ]
    def diagnostic(sid, date, field, quote):
        return {'record_type': 'diagnostic_test', 'source_ids': [sid],
                'diagnostic_result': {'date': date, 'facility': 'Synthetic Imaging',
                    'provider': 'Example, MD', 'study': 'Lumbar MRI',
                    'evidence': [{'source_id': sid, 'date_quote': field, 'quote': quote}]}}
    data = {'sources': [{'id': 'D001', 'scope': 'medical'}, {'id': 'D002', 'scope': 'medical'}],
            'entries': [
                diagnostic('D001', '01/12/2026', 'Collected On 01/12/26 0900', 'Impression: No acute fracture.'),
                {'record_type': 'clinical_care', 'source_ids': ['D001'],
                 'text': '01/12/2026. Synthetic Clinic. Example, MD. Office Visit. Improving pain.'},
                diagnostic('D002', '01/19/2026', 'Date of service: 01/19/2026', 'Impression: No canal stenosis.'),
            ]}
    return docs, data


def test_uncertain_diagnostic_withholds_entire_source_and_remaps_supported_evidence():
    docs, data = fixture()
    before = copy.deepcopy(data)
    with pytest.raises(ScopeReviewRequired):
        parse_scoped_response(json.dumps(data), docs)
    text, exclusions, reviews = parse_draft_response(json.dumps(data), docs)
    assert text == '01/19/2026. Synthetic Imaging. Example, MD. Lumbar MRI. Impression: No canal stenosis.'
    assert not exclusions and len(reviews) == 1
    assert reviews[0]['source_id'] == 'D001' and reviews[0]['source_pages'] == [7]
    assert len(reviews[0]['withheld_proposed_entries']) == 2
    assert reviews[0]['diagnostic_findings'][0]['entry_number'] == 1
    assert 'service_date_not_established' in reviews[0]['diagnostic_findings'][0]['failed_checks']
    assert data == before
    review = review_markdown(reviews)
    assert DRAFT_LABEL in review and 'NOT APPROVED' in review and 'Physical PDF pages' in review
    assert 'Collected On 01/12/26 0900' in review and 'Impression: No acute fracture.' in review
    assert 'Office Visit' not in text


def test_every_rejected_candidate_in_one_source_has_review_details():
    docs, data = fixture()
    second = copy.deepcopy(data['entries'][0])
    second['diagnostic_result']['evidence'][0]['date_quote'] = 'Report Date: 01/12/2026'
    data['entries'].append(second)
    _, _, reviews = parse_draft_response(json.dumps(data), docs)
    assert [f['entry_number'] for f in reviews[0]['diagnostic_findings']] == [1, 4]


@pytest.mark.parametrize('change', [
    'unknown_nested_id', 'unrelated_nested_id', 'missing_evidence', 'missing_header',
    'duplicate_source', 'missing_source', 'unknown_entry_id', 'duplicate_entry_id',
    'conflicting_scope', 'excluded_source', 'invalid_record_type', 'non_string_text',
])
def test_review_cannot_hide_structural_or_source_integrity_errors(change):
    docs, data = fixture()
    entry = data['entries'][-1]  # Must not hide a later error behind the first rejection.
    if change == 'unknown_nested_id':
        entry['diagnostic_result']['evidence'][0]['source_id'] = 'D099'
    elif change == 'unrelated_nested_id':
        entry['diagnostic_result']['evidence'][0]['source_id'] = 'D001'
    elif change == 'missing_evidence':
        entry['diagnostic_result']['evidence'] = []
    elif change == 'missing_header':
        del entry['diagnostic_result']['study']
    elif change == 'duplicate_source':
        data['sources'].append(copy.deepcopy(data['sources'][0]))
    elif change == 'missing_source':
        data['sources'].pop()
    elif change == 'unknown_entry_id':
        entry['source_ids'] = ['D099']
    elif change == 'duplicate_entry_id':
        entry['source_ids'] *= 2
    elif change == 'conflicting_scope':
        data['sources'].append({'id': 'D001', 'scope': 'excluded'})
    elif change == 'excluded_source':
        data['sources'][1].update(scope='excluded', category='correspondence', reason='Synthetic letter')
    elif change == 'invalid_record_type':
        entry['record_type'] = 'invented-type'
    else:
        entry['text'] = 5
    with pytest.raises(ScopeReviewRequired):
        parse_draft_response(json.dumps(data), docs)


def test_linked_sources_are_reviewed_together_without_orphaning_or_publishing_entries():
    docs, data = fixture()
    data['entries'][1]['source_ids'] = ['D001', 'D002']
    text, exclusions, reviews = parse_draft_response(json.dumps(data), docs)
    assert not text and not exclusions
    assert {r['source_id'] for r in reviews} == {'D001', 'D002'}
    assert 'review these source sections together' in reviews[1]['reason']
    assert 'incomplete until reviewed' in review_markdown(reviews)


def test_review_propagates_through_transitive_shared_source_references():
    docs, data = fixture()
    docs.append({'filename': 'third.txt', 'content': 'Synthetic clinical source.'})
    data['sources'].append({'id': 'D003', 'scope': 'medical'})
    clinical = copy.deepcopy(data['entries'][1])
    clinical['source_ids'] = ['D002', 'D003']
    # Visit this link before its neighbor is pending to exercise fixed-point propagation.
    data['entries'].insert(1, clinical)
    data['entries'][2]['source_ids'] = ['D001', 'D002']
    text, exclusions, reviews = parse_draft_response(json.dumps(data), docs)
    assert not text and not exclusions
    assert {r['source_id'] for r in reviews} == {'D001', 'D002', 'D003'}


def test_diagnostic_review_reuses_checkpoint_without_retrying_model(tmp_path):
    docs, data = fixture()
    checkpoint = tmp_path / 'scope.json'
    raw = json.dumps(data)
    calls = []
    def call(*args, **kwargs):
        calls.append(1)
        return raw
    first = screen_batch('synthetic prompt', docs, call, checkpoint=checkpoint, allow_manual_review=True)
    saved = json.loads(checkpoint.read_text())
    assert saved['status'] == 'complete_with_review'
    assert saved['attempts'] == [{'response': raw}]
    assert saved['manual_reviews'][0]['status'] == 'pending'
    assert screen_batch('synthetic prompt', docs, call, checkpoint=checkpoint,
                        allow_manual_review=True) == first
    assert calls == [1]
    assert json.loads(checkpoint.read_text()) == saved


def test_previously_blocked_diagnostic_is_not_automatically_reinterpreted(tmp_path):
    docs, data = fixture()
    checkpoint = tmp_path / 'scope.json'
    with pytest.raises(ScopeReviewRequired):
        screen_batch('same', docs, lambda *a, **k: json.dumps(data), checkpoint=checkpoint)
    before = checkpoint.read_bytes()
    with pytest.raises(ScopeReviewRequired, match='still needs review'):
        screen_batch('same', docs, lambda *a, **k: pytest.fail('No retry'),
                     checkpoint=checkpoint, allow_manual_review=True)
    assert checkpoint.read_bytes() == before


def test_old_protocol_checkpoint_stays_untouched_without_new_model_calls(tmp_path):
    docs, data = fixture()
    checkpoint = tmp_path / 'scope.json'
    signature = hashlib.sha256(json.dumps({'protocol': 1, 'prompt': 'same', 'documents': docs,
                                          'model': None}, sort_keys=True).encode()).hexdigest()
    checkpoint.write_text(json.dumps({'signature': signature, 'status': 'blocked',
                                     'attempts': [{'response': json.dumps(data),
                                                   'code': 'diagnostic_source_review'}]}))
    before = checkpoint.read_bytes()
    with pytest.raises(ScopeReviewRequired, match='changed'):
        screen_batch('same', docs, lambda *a, **k: pytest.fail('No retry'),
                     checkpoint=checkpoint, allow_manual_review=True)
    assert checkpoint.read_bytes() == before


@pytest.mark.parametrize('all_deferred', [False, True])
def test_review_exports_are_incomplete_and_never_mix_candidates_into_chronology(tmp_path, all_deferred):
    docs, data = fixture()
    if all_deferred:
        data['entries'][1]['source_ids'] = ['D001', 'D002']
    a = ChronologyAgent.__new__(ChronologyAgent)
    a.model = 'synthetic'; a.logger = logging.getLogger('synthetic-review-export')
    a._read_extracted_files = lambda _: docs
    a._call_api_with_retry = lambda *a, **k: json.dumps(data)
    batches, output, sources = tmp_path / 'batches', tmp_path / 'output', tmp_path / 'sources'
    sources.mkdir()
    a.generate_batches(str(sources), str(batches))
    a._call_api_with_retry = lambda *a, **k: pytest.fail('No generation or consolidation retry')
    assert a.generate_batches(str(sources), str(batches))['batches_skipped_from_disk'] == 1
    a.extract_header = lambda *a: {'patient_name': 'SYNTHETIC', 'date_of_birth': '[See Records]',
                                 'date_of_injury': '[See Records]'}
    a.generate_summary_and_gaps = lambda *a: {'summary_md': 'Synthetic summary.', 'gaps_md': 'Synthetic gaps.'}
    result = a.assemble_outputs(str(sources), str(batches), str(output))
    assert result['manual_review_count'] == (2 if all_deferred else 1)
    exported = json.loads((output / 'chronology.json').read_text())
    assert exported['manual_review_required']
    assert 'No acute fracture' not in json.dumps(exported['records'])
    review_json = json.loads((output / 'manual_review.json').read_text())
    assert review_json == exported['manual_reviews']
    assert review_json[0]['source_pages'] == [7]
    assert review_json[0]['status'] == 'pending'
    assert review_json[0]['diagnostic_findings']
    for content in [(output / 'chronology.md').read_text(),
                    '\n'.join(p.text for p in Document(output / 'chronology.docx').paragraphs)]:
        assert DRAFT_LABEL in content
        assert ('01/19/2026' in content) == (not all_deferred)
        assert 'No acute fracture' not in content and 'Office Visit' not in content
    review_word = '\n'.join(p.text for p in Document(output / 'manual_review.docx').paragraphs)
    assert 'NOT APPROVED' in review_word and 'No acute fracture' in review_word
    assert 'service_date_not_established' in review_word
