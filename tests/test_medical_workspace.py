"""Medical-only cases use fictional, source-backed fixtures and no external calls."""

import copy
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from src.case_store import CaseStore, MEDICAL_POLICY
from src.medical_evidence import (
    validate_document,
    extract_group,
    exclude_testimony,
    EvidenceReview,
    page_groups,
)
from src.medical_run import MedicalRun, review_source
from src.session_state import SessionState
from src.deposition_evidence import atomic_json

TEXT = """=== SOURCE PDF PAGE 1 ===
Patient: Alex Example
Date of birth: 04/14/1980
Date of service: 02/05/2026
Facility: Example Clinic
Provider: Avery Example, MD
Office evaluation
History: Neck pain without arm weakness.
Examination: Upper extremity strength 5/5.
Impression: Cervical strain.
Plan: Physical therapy and follow-up in four weeks.
"""
PAGES = [{"page": 1, "text": TEXT}]
CASE = {"name": "Alex Example", "dob": "04/14/1980"}
DOC = {"id": "doc-one", "path": "record.pdf", "sha256": "source-sha"}


def response():
    return {
        "sections": [
            {
                "pages": [1],
                "scope": "medical",
                "reason": "Clinical encounter",
                "patient": {
                    "name": "Alex Example",
                    "dob": "04/14/1980",
                    "evidence": [
                        {
                            "page": 1,
                            "quote": "Patient: Alex Example\nDate of birth: 04/14/1980",
                        }
                    ],
                },
                "entries": [
                    {
                        "record_type": "clinical_care",
                        "service_dates": ["02/05/2026"],
                        "date_evidence": [
                            {
                                "date": "02/05/2026",
                                "page": 1,
                                "quote": "Date of service: 02/05/2026",
                            }
                        ],
                        "facility": "Example Clinic",
                        "provider": "Avery Example, MD",
                        "service_name": "Office evaluation",
                        "body": "History: Neck pain without arm weakness. Examination: Upper extremity strength 5/5. Impression: Cervical strain. Plan: Physical therapy and follow-up in four weeks.",
                        "evidence": [{"page": 1, "quote": TEXT}],
                    }
                ],
            }
        ]
    }


def test_entries_retain_original_page_and_source_version():
    result = validate_document(response(), PAGES, CASE, MEDICAL_POLICY, DOC)
    entry = result["entries"][0]
    assert (
        entry["evidence"][0]["page"] == 1
        and entry["evidence"][0]["source_sha256"] == "source-sha"
    )
    assert entry["service_dates"] == ["02/05/2026"]
    assert entry["text"].startswith(
        "02/05/2026. Example Clinic. Avery Example, MD. Office evaluation."
    )
    assert "Visit Type:" not in entry["text"]


def test_wrong_patient_is_withheld_without_altering_source():
    result = validate_document(
        response(),
        PAGES,
        {"name": "Wrong Patient", "dob": "04/14/1980"},
        MEDICAL_POLICY,
        DOC,
    )
    assert not result["entries"] and result["reviews"][0]["kind"] == "patient_identity"
    assert PAGES[0]["text"] == TEXT


def test_unquoted_claim_is_not_accepted_as_evidence():
    data = response()
    data["sections"][0]["entries"][0]["evidence"][0][
        "quote"
    ] = "Invented surgery was performed."
    with pytest.raises(EvidenceReview, match="does not match"):
        validate_document(data, PAGES, CASE, MEDICAL_POLICY, DOC)


def test_birth_date_cannot_be_substituted_for_service_date():
    data = response()
    record = data["sections"][0]["entries"][0]
    record["service_dates"] = ["04/14/1980"]
    record["date_evidence"] = [
        {"date": "04/14/1980", "page": 1, "quote": "Date of birth: 04/14/1980"}
    ]
    with pytest.raises(EvidenceReview, match="service date"):
        validate_document(data, PAGES, CASE, MEDICAL_POLICY, DOC)


def test_deposition_preface_excluded_clinical_attachment_preserved():
    pages = [
        {
            "page": 1,
            "text": "AI-GENERATED DEPOSITION SUMMARY\nSummary of the deposition. Neck pain was discussed.",
        },
        {
            "page": 2,
            "text": "DEPOSITION OF FICTIONAL WITNESS\nQ. State your name.\nA. Example.\nQ. Where were you?\nA. Home.",
        },
        {"page": 3, "text": TEXT.replace("PAGE 1", "PAGE 3")},
    ]
    retained, excluded = exclude_testimony(pages)
    assert [p["page"] for p in retained] == [3]
    assert {p for x in excluded for p in x["pages"]} == {1, 2}


def test_legal_and_billing_inventory_never_becomes_clinical_entry():
    data = {
        "sections": [
            {
                "pages": [1],
                "scope": "excluded",
                "category": "legal_document",
                "reason": "Lien agreement, not medical care.",
            }
        ]
    }
    result = validate_document(data, PAGES, CASE, MEDICAL_POLICY, DOC)
    assert (
        not result["entries"] and result["excluded"][0]["category"] == "legal_document"
    )
    data = response()
    data["sections"][0]["entries"][0]["record_type"] = "medical_billing"
    result = validate_document(data, PAGES, CASE, MEDICAL_POLICY, DOC)
    assert not result["entries"] and result["excluded"][0]["category"] == "billing_only"


def test_supported_medical_expert_report_is_included():
    data = response()
    data["sections"][0]["entries"][0]["record_type"] = "medical_evaluation"
    assert (
        validate_document(data, PAGES, CASE, MEDICAL_POLICY, DOC)["entries"][0][
            "record_type"
        ]
        == "medical_evaluation"
    )


def test_complete_page_groups_do_not_truncate_long_sources():
    pages = [{"page": i, "text": str(i) + "x" * 20000} for i in range(1, 10)]
    assert [p for group in page_groups(pages) for p in group] == pages


def model_call(prompt, max_tokens):
    if prompt.startswith("Check the companion"):
        return json.dumps(
            {
                "supported": True,
                "complete": True,
                "reason": "The fictional overview matches the supplied entries.",
            }
        )
    if prompt.startswith("Prepare an executive"):
        records = json.loads(prompt[prompt.index("\n[") + 1 :])
        return json.dumps(
            {
                "summary": "The supplied records document evaluation and conservative treatment of cervical symptoms.",
                "gaps": "No candidate gaps identified in these supplied entries.",
                "entry_ids": [e["id"] for e in records],
            }
        )
    if prompt.startswith("Verify the candidate"):
        payload = json.loads(prompt[prompt.index('{"policy"') :])
        return json.dumps(
            {
                "entries": [
                    {
                        "id": e["id"],
                        "verdict": "supported",
                        "reason": "All clauses match exact source text.",
                    }
                    for e in payload["candidate"]["entries"]
                ],
                "missing_encounters": [],
            }
        )
    return json.dumps(response())


def test_cached_source_checks_survive_restart_without_more_calls(tmp_path):
    first = extract_group(
        PAGES, CASE, MEDICAL_POLICY, DOC, model_call, tmp_path / "work.json"
    )
    second = extract_group(
        PAGES,
        CASE,
        MEDICAL_POLICY,
        DOC,
        lambda *a, **k: pytest.fail("No repeat model call"),
        tmp_path / "work.json",
    )
    assert first == second and first["entries"][0]["verification"] == "source_checked"


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setenv("SESSION_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("REQUIRE_PERSISTENT_STORAGE", "false")
    return CaseStore(tmp_path)


def create_case(store):
    def create(source, name):
        return store.sessions.create(
            "synthetic-case",
            name,
            source,
            "/Medical chronology pipeline outputs/synthetic-case",
        )

    return store.create(
        SimpleNamespace(create_session=create),
        name=CASE["name"],
        dob=CASE["dob"],
        doi="02/03/2026",
        source="https://www.dropbox.com/fictional",
        model="synthetic-model",
    )


def test_repeated_resume_is_one_job_and_restart_recovers_lease(store):
    case = create_case(store)
    a = store.enqueue(case["id"])
    b = store.enqueue(case["id"])
    assert a["id"] == b["id"]
    job = store.claim("old-worker")
    assert job["id"] == a["id"]
    assert store.claim("new-worker") is None
    with store.connect() as db:
        db.execute("UPDATE jobs SET heartbeat=0")
    resumed = store.claim("new-worker")
    assert resumed["id"] == job["id"]


def test_archive_restore_preserves_original_tree(store):
    case = create_case(store)
    p = store.sessions.input_dir(case["id"]) / "source.pdf"
    p.write_bytes(b"original")
    store.archive(case["id"])
    assert store.list_cases() == []
    assert store.list_cases(True)[0]["id"] == case["id"]
    store.archive(case["id"], True)
    assert p.read_bytes() == b"original"
    assert store.overview(case["id"])["policy"] == MEDICAL_POLICY


def test_corrupt_case_is_visible_and_reads_do_not_create_cases(store):
    path = store.sessions.sessions_root / "broken"
    path.mkdir()
    (path / "state.json").write_text("bad")
    assert store.list_cases()[0]["status"] == "recovery_required"
    with pytest.raises(FileNotFoundError):
        store.overview("does-not-exist")
    assert not (store.sessions.sessions_root / "does-not-exist").exists()


def test_case_scoped_file_routes_refuse_traversal(store):
    case = create_case(store)
    with pytest.raises(FileNotFoundError):
        store.file(case["id"], "input", "../../workspace.sqlite3")


def test_document_update_does_not_erase_other_records(store):
    create_case(store)
    for number in (1, 2):
        d = {"id": f"d{number}", "sha256": str(number)}
        store.replace_document_result(
            "synthetic-case", d, [{"id": f"e{number}", "document_id": d["id"]}], []
        )
    store.replace_document_result(
        "synthetic-case",
        {"id": "d1", "sha256": "new"},
        [{"id": "new-e1", "document_id": "d1"}],
        [],
    )
    assert {e["id"] for e in store.all("synthetic-case", "entries")} == {"new-e1", "e2"}


def prepared_pipeline(store):
    from src.pipeline import MedicalChronologyPipeline

    case = create_case(store)
    sid = case["id"]
    root = store.sessions.input_dir(sid)
    pdf = root / "record.pdf"
    pdf.write_bytes(b"fictional-source-pdf-bytes")
    state = store.sessions.load(sid)
    state.phases["download"].status = "complete"
    state.phases["download"].data = {
        "manifest_version": 1,
        "manifest": [
            {
                "path": "record.pdf",
                "size": pdf.stat().st_size,
                "sha256": hashlib.sha256(pdf.read_bytes()).hexdigest(),
            }
        ],
    }
    store.sessions.save(state)
    p = MedicalChronologyPipeline.__new__(MedicalChronologyPipeline)
    p.store = store.sessions
    calls = {"ocr": 0, "model": 0, "upload": 0}

    def extract(path, **kwargs):
        calls["ocr"] += 1
        return {
            "success": True,
            "text": TEXT,
            "page_count": 1,
            "page_results": [{"page": 1, "status": "text"}],
            "source_path": path,
            "file_name": "record.pdf",
        }

    def save(result, output, input_root):
        target = Path(output) / "record.txt"
        target.parent.mkdir(exist_ok=True)
        target.write_text(result["text"])
        return str(target)

    def model(prompt, **kwargs):
        calls["model"] += 1
        return model_call(prompt, **kwargs)

    def upload(**kwargs):
        calls["upload"] += 1
        return {
            "success": True,
            "uploaded": [
                {"name": f.name, "verified": True}
                for f in Path(kwargs["local_dir"]).iterdir()
                if f.is_file()
            ],
        }

    p.ocr_client = SimpleNamespace(extract_text=extract, save_extracted_text=save)
    p.chronology_agent = SimpleNamespace(
        model="synthetic-model", _call_api_with_retry=model
    )
    p.dropbox_tool = SimpleNamespace(upload_folder=upload)
    return case, p, calls


def test_full_medical_job_exports_and_resume_reuses_unchanged_sources(store):
    from src.workspace_worker import execute_job

    case, p, calls = prepared_pipeline(store)
    store.enqueue(case["id"])
    job = store.claim("worker-one")
    execute_job(store, job, lambda model: p)
    assert store.job(case["id"])["status"] == "complete"
    final = store.overview(case["id"])
    assert final["status"] == "complete"
    assert final["phases"]["upload"]["data"]["verified_all"] is True
    artifacts = store.all(case["id"], "artifacts")
    assert {
        "chronology.docx",
        "chronology.json",
        "summary.md",
        "gaps.md",
        "manifest.json",
    } <= {a["name"] for a in artifacts}
    hashes = {a["id"]: a["sha256"] for a in artifacts}
    before = dict(calls)
    store.enqueue(case["id"])
    execute_job(store, store.claim("worker-two"), lambda model: p)
    assert store.job(case["id"])["status"] == "complete"
    assert calls["ocr"] == before["ocr"] and calls["model"] == before["model"]
    assert {a["id"]: a["sha256"] for a in store.all(case["id"], "artifacts")} == hashes


def test_targeted_retry_retains_unrelated_ocr_and_entries(store):
    case, p, calls = prepared_pipeline(store)
    MedicalRun(store, p, case["id"]).run()
    doc = store.all(case["id"], "documents")[0]
    other = {
        "id": "other",
        "path": "other.pdf",
        "sha256": "other-sha",
        "status": "included",
    }
    store.replace_document_result(
        case["id"],
        other,
        [
            {
                "id": "other-entry",
                "document_id": "other",
                "text": "Unrelated retained entry",
            }
        ],
        [],
    )
    from src.medical_run import issue_for

    issue = issue_for(doc, "ocr", "Retry this page", [1])
    store.put(case["id"], "issues", issue["id"], issue)
    review_source(
        store,
        case["id"],
        {
            "target": issue["id"],
            "fingerprint": issue["fingerprint"],
            "action": "retry",
            "reviewer": "Test Reviewer",
            "reason": "Checking the original source page.",
        },
    )
    assert (
        store.get(case["id"], "entries", "other-entry")["text"]
        == "Unrelated retained entry"
    )
    assert (store.directory(case["id"]) / "batches" / "run_model.json").exists()
    assert store.history(case["id"])[0]["action"] == "retry"


def test_review_rejects_stale_fingerprint_and_unnamed_resolution(store):
    case, p, calls = prepared_pipeline(store)
    MedicalRun(store, p, case["id"]).run()
    doc = store.all(case["id"], "documents")[0]
    with pytest.raises(ValueError, match="changed"):
        review_source(
            store,
            case["id"],
            {
                "target": doc["id"],
                "fingerprint": "old",
                "action": "defer",
                "reviewer": "Reviewer",
                "reason": "Source review",
            },
        )
    with pytest.raises(ValueError, match="name"):
        review_source(
            store,
            case["id"],
            {
                "target": doc["id"],
                "fingerprint": doc["sha256"],
                "action": "defer",
                "reviewer": "",
                "reason": "Source review",
            },
        )


def test_clinical_exhibit_continuation_is_not_discarded_by_transcript_filter():
    continuation = {
        "page": 3,
        "text": "Assessment continued: improving function. Return in six weeks.",
    }
    pages = [
        {"page": 1, "text": "DEPOSITION OF FICTIONAL WITNESS\nQ. Name?\nA. Example."},
        {"page": 2, "text": TEXT.replace("PAGE 1", "PAGE 2")},
        continuation,
    ]
    retained, excluded = exclude_testimony(pages)
    assert continuation in retained and {p["page"] for p in retained} == {2, 3}


def test_identity_birth_date_must_have_birth_date_role():
    data = response()
    data["sections"][0]["patient"]["evidence"] = [
        {"page": 1, "quote": "Patient: Alex Example"},
        {"page": 1, "quote": "Date of service: 02/05/2026"},
    ]
    data["sections"][0]["patient"]["dob"] = "02/05/2026"
    result = validate_document(data, PAGES, {"name": CASE["name"]}, MEDICAL_POLICY, DOC)
    assert not result["entries"] and result["reviews"][0]["kind"] == "patient_identity"


@pytest.mark.parametrize(
    "birth_field",
    [
        "Date of Birth: Apr 14, 1980",
        "Date of birth: April 14, 1980",
        "DOB: 1980-04-14",
        "Birth date: 04-14-1980",
        "Date of birth\n: Apr 14, 1980",
    ],
)
@pytest.mark.parametrize("case_dob", ["", "04/14/1980"])
def test_birth_date_presentation_does_not_withhold_documented_visit(
    birth_field, case_dob
):
    data = response()
    text = TEXT.replace("Date of birth: 04/14/1980", birth_field)
    patient = data["sections"][0]["patient"]
    patient["evidence"][0]["quote"] = "Patient: Alex Example\n" + birth_field
    data["sections"][0]["entries"][0]["evidence"][0]["quote"] = text
    result = validate_document(
        data,
        [{"page": 1, "text": text}],
        {"name": CASE["name"], "dob": case_dob},
        MEDICAL_POLICY,
        DOC,
    )
    assert len(result["entries"]) == 1
    assert not result["reviews"]
    assert result["entries"][0]["source_patient"]["dob"] == "04/14/1980"
    # The source quotation remains verbatim, with the original month/date format.
    assert birth_field in result["entries"][0]["evidence"][0]["quote"]


@pytest.mark.parametrize(
    "birth_field",
    [
        "DOB: Apr 14, 1981",  # Actual discrepancy, not a format difference.
        "DOB: 04/14/80",  # Do not infer a century.
        "DOB: Apr 14, 19800",  # Do not accept a valid prefix of an invalid year.
        "DOB: Feb 30, 1980",  # Invalid calendar date.
        "DOB: unknown\nDate of service: Apr 14, 1980",
        "Age: 46\nDate of visit: Apr 14, 1980",
    ],
)
def test_birth_date_format_support_does_not_accept_unverified_identity(birth_field):
    data = response()
    text = TEXT.replace("Date of birth: 04/14/1980", birth_field)
    data["sections"][0]["patient"]["evidence"][0]["quote"] = (
        "Patient: Alex Example\n" + birth_field
    )
    data["sections"][0]["entries"][0]["evidence"][0]["quote"] = text
    result = validate_document(
        data, [{"page": 1, "text": text}], CASE, MEDICAL_POLICY, DOC
    )
    assert not result["entries"]
    assert result["reviews"][0]["kind"] == "patient_identity"


def test_written_month_conflict_is_detected_outside_selected_identity_quote():
    data = response()
    text = TEXT + "\nPatient: Alex Example\nDOB: April 14, 1981\n"
    data["sections"][0]["entries"][0]["evidence"][0]["quote"] = text
    result = validate_document(
        data, [{"page": 1, "text": text}], CASE, MEDICAL_POLICY, DOC
    )
    assert not result["entries"]
    assert result["reviews"][0]["kind"] == "patient_identity"


def formatted_encounter_response():
    data = response()
    text = TEXT.replace("Date of birth: 04/14/1980", "Date of Birth: Apr 14, 1980")
    text = text.replace(
        "Date of service: 02/05/2026", "Date of Visit: February 05, 2026"
    )
    signature = (
        "Electronically Signed By Avery Example, MD on February 05, 2026 08:38 AM"
    )
    cosigner = "Electronically Co-Signed By Robin Example, MD"
    text += "\n" + signature + "\n" + cosigner + " on February 07, 2026\n"
    section = data["sections"][0]
    section["patient"]["evidence"] = [
        {"page": 1, "quote": "Patient: Alex Example\nDate of Birth: Apr 14, 1980"}
    ]
    record = section["entries"][0]
    record["provider"] = "Avery Example, MD; " + cosigner
    record["date_evidence"] = [
        {"date": "02/05/2026", "page": 1, "quote": "Date of Visit: February 05, 2026"},
        {"date": "02/05/2026", "page": 1, "quote": signature},
    ]
    record["evidence"] = [{"page": 1, "quote": text}]
    return data, [{"page": 1, "text": text}]


def test_formatted_visit_reaches_independent_audit_with_roles_preserved(tmp_path):
    data, pages = formatted_encounter_response()
    audits = []

    def call(prompt, max_tokens):
        if prompt.startswith("Verify the candidate"):
            payload = json.loads(prompt[prompt.index('{"policy"') :])
            audits.append(payload["candidate"]["entries"])
            return model_call(prompt, max_tokens)
        return json.dumps(data)

    result = extract_group(
        pages, {"name": CASE["name"]}, MEDICAL_POLICY, DOC, call, tmp_path / "work.json"
    )
    assert len(result["entries"]) == 1 and not result["reviews"]
    assert len(audits) == 1 and len(audits[0]) == 1
    entry = result["entries"][0]
    assert entry["service_dates"] == ["02/05/2026"]
    assert "Co-Signed By Robin Example, MD" in entry["provider"]
    assert entry["verification"] == "source_checked"


@pytest.mark.parametrize(
    "failure",
    [
        "signature_only",
        "conflicting_service",
        "invented_signature",
        "invented_provider",
    ],
)
def test_formatted_visit_still_rejects_unsupported_dates_or_provider(failure):
    data, pages = formatted_encounter_response()
    record = data["sections"][0]["entries"][0]
    if failure == "signature_only":
        record["date_evidence"] = record["date_evidence"][1:]
    elif failure == "conflicting_service":
        pages[0]["text"] += "\nDate of Visit: February 06, 2026\n"
        record["date_evidence"].append(
            {
                "date": "02/05/2026",
                "page": 1,
                "quote": "Date of Visit: February 06, 2026",
            }
        )
    elif failure == "invented_signature":
        record["date_evidence"][1][
            "quote"
        ] = "Electronically Signed By Invented Clinician on February 05, 2026"
    else:
        record["provider"] += "; Invented Clinician, MD"
    with pytest.raises(EvidenceReview):
        validate_document(data, pages, CASE, MEDICAL_POLICY, DOC)


def test_historical_and_new_cases_are_isolated_and_collisions_preserved(store):
    from src.session_state import SessionStore

    old = SessionStore(str(store.root.parent))
    old.create(
        "historical",
        "History",
        "https://www.dropbox.com/fictional",
        "/Medical chronology pipeline outputs/historical",
    )
    before = (old.session_dir("historical") / "state.json").read_bytes()
    create_case(store)
    assert {c["id"] for c in store.list_cases()} == {"synthetic-case", "historical"}
    assert [s.session_id for s in old.list_sessions()] == ["historical"]
    with pytest.raises(ValueError, match="Historical"):
        store.enqueue("historical")
    assert (old.session_dir("historical") / "state.json").read_bytes() == before
    assert not (store.root / "medical-sessions" / "historical").exists()
    (store.root / "medical-sessions" / "historical").mkdir()
    with pytest.raises(ValueError, match="conflicts"):
        store.directory("historical")


def test_newer_database_schema_fails_without_rewriting_metadata(store):
    with store.connect() as db:
        db.execute("PRAGMA user_version=99")
    with pytest.raises(RuntimeError, match="newer version"):
        CaseStore(store.root.parent)
    with store.connect() as db:
        assert db.execute("PRAGMA user_version").fetchone()[0] == 99


def test_source_root_symlink_cannot_escape_case(store, tmp_path):
    case = create_case(store)
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "secret.pdf").write_bytes(b"outside")
    root = store.directory(case["id"]) / "input"
    root.symlink_to(outside, target_is_directory=True)
    with pytest.raises(FileNotFoundError, match="outside"):
        store.file(case["id"], "input", "secret.pdf")


def test_export_after_review_requires_targeted_processing_first(store):
    case, p, calls = prepared_pipeline(store)
    MedicalRun(store, p, case["id"]).run()
    doc = store.all(case["id"], "documents")[0]
    decision = {
        "target": doc["id"],
        "fingerprint": doc["sha256"],
        "action": "defer",
        "reviewer": "Test Reviewer",
        "reason": "Inspect this source.",
    }
    review_source(store, case["id"], decision)
    before = calls["upload"]
    with pytest.raises(ValueError, match="Resume document"):
        MedicalRun(store, p, case["id"]).run("export")
    assert calls["upload"] == before
    with pytest.raises(ValueError, match="changed"):
        review_source(store, case["id"], decision)


def test_changed_original_blocks_export_and_retains_previous_version(store):
    case, p, calls = prepared_pipeline(store)
    MedicalRun(store, p, case["id"]).run()
    versions = store.all(case["id"], "artifacts")
    before = calls["upload"]
    (store.directory(case["id"]) / "input" / "record.pdf").write_bytes(
        b"changed original"
    )
    with pytest.raises(Exception):
        MedicalRun(store, p, case["id"]).run("export")
    assert calls["upload"] == before and store.all(case["id"], "artifacts") == versions


def test_verified_subset_is_not_successful_delivery(store):
    case, p, calls = prepared_pipeline(store)
    p.dropbox_tool.upload_folder = lambda **kw: {
        "success": True,
        "uploaded": [{"name": "chronology.docx", "verified": True}],
    }
    with pytest.raises(RuntimeError, match="every file"):
        MedicalRun(store, p, case["id"]).run()
    assert store.overview(case["id"])["phases"]["upload"]["status"] != "complete"
    assert any(
        a["name"] == "chronology.docx" for a in store.all(case["id"], "artifacts")
    )


def test_companion_failure_keeps_detailed_flagged_draft_available(store):
    case, p, calls = prepared_pipeline(store)

    def model(prompt, **kw):
        if prompt.startswith("Check the companion"):
            return json.dumps(
                {
                    "supported": False,
                    "complete": False,
                    "reason": "The overview added an unsupported finding.",
                }
            )
        return model_call(prompt, **kw)

    p.chronology_agent._call_api_with_retry = model
    MedicalRun(store, p, case["id"]).run()
    case = store.overview(case["id"])
    assert case["status"] == "complete"
    assert (
        case["phases"]["header"]["status"] == "complete"
        and case["phases"]["summary"]["status"] == "failed"
    )
    assert case["review_count"] == 1
    folder = store.directory(case["id"]) / "output"
    assert (
        "Draft complete—companion report review required"
        in (folder / "chronology.md").read_text()
    )
    assert (
        "Unresolved source sections are withheld"
        not in (folder / "chronology.md").read_text()
    )
    assert "unsupported finding" in (folder / "manual_review.md").read_text()
    assert "unavailable" in (folder / "summary.md").read_text()
    assert json.loads((folder / "chronology.json").read_text())["records"]


def test_partial_document_questions_do_not_prevent_flagged_delivery(store):
    case, p, calls = prepared_pipeline(store)

    def model(prompt, **kw):
        if prompt.startswith("Verify the candidate"):
            payload = json.loads(prompt[prompt.index('{"policy"') :])
            return json.dumps(
                {
                    "entries": [
                        {
                            "id": e["id"],
                            "verdict": "uncertain",
                            "reason": "Original procedure level requires team review.",
                        }
                        for e in payload["candidate"]["entries"]
                    ],
                    "missing_encounters": [],
                }
            )
        return model_call(prompt, **kw)

    p.chronology_agent._call_api_with_retry = model
    MedicalRun(store, p, case["id"]).run()
    assert store.overview(case["id"])["status"] == "complete"
    payload = json.loads(
        (store.directory(case["id"]) / "output" / "chronology.json").read_text()
    )
    assert payload["records"] == [] and payload["manual_reviews"]
    text = (store.directory(case["id"]) / "output" / "manual_review.md").read_text()
    assert "record.pdf" in text and "PDF pages 1" in text


def test_complete_companion_coverage_across_large_cases(tmp_path):
    from src.companion_reports import generate_reports

    entries = [
        {
            "id": str(i),
            "text": "A fictional long source-backed encounter. " * 450,
            "service_dates": ["02/05/2026"],
        }
        for i in range(9)
    ]
    files = generate_reports(entries, model_call, tmp_path)
    coverage = json.loads(files["companion_coverage.json"])
    assert (
        coverage["entry_ids"] == [str(i) for i in range(9)] and coverage["groups"] > 1
    )


def test_active_case_lock_preserves_state_when_duplicate_worker_attempts(store):
    from src.workspace_worker import execute_job

    case, p, calls = prepared_pipeline(store)
    store.enqueue(case["id"])
    job = store.claim("first")
    before = (store.directory(case["id"]) / "state.json").read_bytes()
    with store.lock(case["id"]):
        execute_job(
            store, job, lambda model: pytest.fail("Must not start a second pipeline")
        )
    assert (store.directory(case["id"]) / "state.json").read_bytes() == before
    assert store.job(case["id"])["status"] == "paused"


@pytest.mark.parametrize(
    "merge_mode", ["valid", "format", "wrong_provider", "wrong_date"]
)
def test_same_day_procedures_merge_all_sources_but_keep_other_provider(
    store, merge_mode
):
    case, p, calls = prepared_pipeline(store)
    run = MedicalRun(store, p, case["id"])
    doc = run.inventory()[0]
    entry = validate_document(response(), PAGES, CASE, MEDICAL_POLICY, doc)["entries"][
        0
    ]
    second = {
        **copy.deepcopy(entry),
        "id": "second-procedure",
        "provider": "Avery Example, M.D.",
        "text": "02/05/2026. Example Clinic. Avery Example, M.D. Cervical injection. History: Persistent symptoms. Procedure: Left C5-6 injection, tolerated without immediate complication.",
    }
    second["evidence"] = [
        {
            **entry["evidence"][0],
            "quote": "Left C5-6 injection, tolerated without immediate complication.",
        }
    ]
    third = {
        **copy.deepcopy(entry),
        "id": "other-provider",
        "provider": "Riley Example, DPT",
    }
    third["text"] = third["text"].replace("Avery Example, MD", "Riley Example, DPT")
    store.replace_document_result(case["id"], doc, [entry, second, third], [])
    merged_text = "02/05/2026. Example Clinic. Avery Example, MD. Evaluation and cervical injection. History: Neck pain without arm weakness. Examination: Upper extremity strength 5/5. Impression: Cervical strain. Procedure: Left C5-6 injection, tolerated without immediate complication. Plan: Physical therapy and follow-up in four weeks."

    merge_calls = []

    def model(prompt, **kw):
        if prompt.startswith("Consolidate these source-checked"):
            inputs = json.loads(
                prompt[prompt.index("\n[") + 1 :].split("\nFORMAT CORRECTION:")[0]
            )
            assert len(inputs) == 2
            merge_calls.append(prompt)
            proposed = merged_text
            if merge_mode == "format" and len(merge_calls) == 1:
                proposed = merged_text.replace(
                    "02/05/2026. Example Clinic. Avery Example, MD.",
                    "02/05/2026 — Avery Example, MD, Example Clinic:",
                )
            if merge_mode == "wrong_provider":
                proposed = merged_text.replace(
                    "Avery Example, MD", "Jordan Different, MD"
                )
            if merge_mode == "wrong_date":
                proposed = merged_text.replace("02/05/2026.", "02/06/2026.", 1)
            return json.dumps(
                {"text": proposed, "covered_ids": [e["id"] for e in inputs]}
            )
        if prompt.startswith("Check this merged"):
            return json.dumps(
                {
                    "supported": True,
                    "complete": True,
                    "reason": "Both clinical records and the distinct injection details are retained.",
                }
            )
        pytest.fail("Unexpected model call")

    p.chronology_agent._call_api_with_retry = model
    result = run.merged_entries()
    assert len(merge_calls) == (2 if merge_mode == "format" else 1)
    if merge_mode.startswith("wrong_"):
        assert len(result) == 1 and result[0]["id"] == third["id"]
        assert any(
            i["kind"] == "same_day_merge" for i in store.all(case["id"], "issues")
        )
        return
    assert len(result) == 2
    merged = next(e for e in result if "merged_entry_ids" in e)
    assert set(merged["merged_entry_ids"]) == {entry["id"], second["id"]}
    assert len(merged["evidence"]) == 2 and "Left C5-6 injection" in merged["text"]
    p.chronology_agent._call_api_with_retry = lambda *a, **k: pytest.fail(
        "Cached merge must survive restart"
    )
    assert MedicalRun(store, p, case["id"]).merged_entries() == result


def test_grouped_therapy_dates_all_remain_source_checked():
    data = response()
    record = data["sections"][0]["entries"][0]
    source = TEXT + "\nDate of service: 02/09/2026\nTreatment: Continued exercises."
    record["service_dates"] = ["02/05/2026", "02/09/2026"]
    record["date_evidence"].append(
        {"date": "02/09/2026", "page": 1, "quote": "Date of service: 02/09/2026"}
    )
    result = validate_document(
        data, [{"page": 1, "text": source}], CASE, MEDICAL_POLICY, DOC
    )
    assert result["entries"][0]["service_dates"] == ["02/05/2026", "02/09/2026"]
    assert "Service dates: 02/05/2026, 02/09/2026." in result["entries"][0]["text"]
    record["date_evidence"].pop()
    with pytest.raises(ValueError, match="every service date"):
        validate_document(
            data, [{"page": 1, "text": source}], CASE, MEDICAL_POLICY, DOC
        )


def test_exact_file_duplicate_preserves_inventory_alias(store):
    case, p, calls = prepared_pipeline(store)
    root = store.directory(case["id"]) / "input"
    duplicate = root / "copy.pdf"
    duplicate.write_bytes((root / "record.pdf").read_bytes())
    state = store.sessions.load(case["id"])
    item = {**state.phases["download"].data["manifest"][0], "path": "copy.pdf"}
    state.phases["download"].data["manifest"].append(item)
    store.sessions.save(state)
    docs = MedicalRun(store, p, case["id"]).inventory()
    assert len(docs) == 2 and sum(d.get("duplicate_kind") == "exact" for d in docs) == 1
    assert duplicate.read_bytes() == (root / "record.pdf").read_bytes()
    assert (
        next(d for d in docs if d.get("duplicate_kind") == "exact")["exclusions"][0][
            "category"
        ]
        == "exact_duplicate"
    )


def test_word_profiles_preserve_narrative_content_and_page_geometry():
    import io
    from docx import Document
    from src.word_export import chronology_docx

    text = (
        "MEDICAL RECORDS SUMMARY\nALEX FICTIONAL\nDate of Birth: 04/14/1980\nDate of Injury: 02/03/2026\n\n"
        + (
            "02/05/2026. Example Clinic. Avery Example, MD. Evaluation. History: "
            + ("A long synthetic detail. " * 100)
        )
    )
    for template in ("classic", "underlined", "plain"):
        d = Document(io.BytesIO(chronology_docx(text, template=template)))
        assert len(d.tables) == 0 and len(d.sections) == 1
        assert (
            d.sections[0].page_width.inches == 8.5
            and d.sections[0].left_margin.inches == 1
        )
        assert (
            d.styles["Normal"].font.name == "Times New Roman"
            and d.styles["Normal"].font.size.pt == 12
        )
        title = next(p for p in d.paragraphs if p.text == "MEDICAL RECORDS SUMMARY")
        assert bool(title.runs[0].underline) == (template == "underlined")
        assert (
            sum(p.text.count("A long synthetic detail.") for p in d.paragraphs) == 100
        )
        assert "PAGE" in d.sections[0].footer._element.xml


def test_rollout_preflight_detects_active_work_without_changing_it(store):
    from scripts.workspace_preflight import inspect

    case = create_case(store)
    before = (store.directory(case["id"]) / "state.json").read_bytes()
    with store.lock(case["id"]):
        result = inspect(store.root)
        assert (
            result["active_case_locks"] == [case["id"]]
            and result["ready_at_check"] is False
        )
    assert inspect(store.root)["active_case_locks"] == []
    store.enqueue(case["id"])
    assert inspect(store.root)["active_jobs"] == [
        {"case_id": case["id"], "status": "queued"}
    ]
    assert (store.directory(case["id"]) / "state.json").read_bytes() == before


def test_real_world_header_shapes_do_not_misread_dob_or_middle_name():
    from src.medical_evidence import same_patient, supports_encounter_date

    header = "EXAMPLE, ALEX MORGAN DOB: 04/14/1980 (46 yo) Acc No. SYNTHETIC DOS:\n03/20/2026"
    assert same_patient("EXAMPLE, ALEX MORGAN", "Alex Morgan Example")
    assert same_patient("EXAMPLE ALEX MORGAN", "Alex Example")
    assert not same_patient("Alex Blake Example", "Alex Morgan Example")
    assert supports_encounter_date("DOS:\n03/20/2026", "03/20/2026", header)
    assert not supports_encounter_date("DOB: 04/14/1980", "04/14/1980", header)
    assert not supports_encounter_date("DOS:\n03/20/2026", "03/20/2026", header + "9")


def test_incomplete_response_is_retained_and_not_retried_unchanged(tmp_path):
    from src.response_recovery import IncompleteResponseError
    from src.medical_evidence import retained_call

    calls = []

    def incomplete(*a, **kw):
        calls.append(1)
        raise IncompleteResponseError("synthetic", "max_tokens", 16000)

    for _ in range(2):
        with pytest.raises(EvidenceReview, match="incomplete response"):
            retained_call(
                tmp_path / "work.json", "same source", incomplete, lambda value: value
            )
    assert len(calls) == 1
    assert json.loads((tmp_path / "work.json").read_text())["status"] == "needs_review"


@pytest.mark.parametrize("label", ["Dates of service", "Service dates", "Visit dates"])
def test_grouped_therapy_accepts_discrete_dates_in_one_explicit_field(label):
    dates = ["03/02/2026", "03/04/2026", "03/09/2026"]
    field = label + ": " + ", ".join(dates) + "."
    text = TEXT.replace("Date of service: 02/05/2026", field)
    data = response()
    entry = data["sections"][0]["entries"][0]
    entry["service_dates"] = dates
    entry["date_evidence"] = [
        {"date": date, "page": 1, "quote": field} for date in dates
    ]
    entry["evidence"] = [{"page": 1, "quote": text}]
    result = validate_document(
        data, [{"page": 1, "text": text}], CASE, MEDICAL_POLICY, DOC
    )
    assert len(result["entries"]) == 1 and not result["reviews"]
    assert result["entries"][0]["service_dates"] == dates
    assert all(date in result["entries"][0]["text"] for date in dates)


@pytest.mark.parametrize(
    "field,date",
    [
        ("Dates of service: 03/02/2026 through 03/09/2026", "03/04/2026"),
        ("Dates of service: 03/02/2026, 03/09/2026", "03/04/2026"),
        ("Dates of service: 03/02/2026, 02/30/2026", "03/02/2026"),
        ("Billing period: 03/02/2026, 03/04/2026", "03/04/2026"),
        ("History: 03/02/2026, 03/04/2026", "03/04/2026"),
    ],
)
def test_attendance_list_does_not_create_dates_from_ranges_or_other_fields(field, date):
    from src.medical_evidence import supports_encounter_date

    assert not supports_encounter_date(field, date, field)


@pytest.mark.parametrize("quoted_role", ["dos", "appointment", "progress"])
def test_emr_date_header_needs_independent_same_page_corroboration(quoted_role):
    from src.medical_evidence import supports_encounter_date

    fields = {
        "dos": "EXAMPLE, ALEX DOB: 04/14/1980 Acc No. TEST DOS:\nSpine & Orthopedic\nSpecialists\n03/20/2026",
        "appointment": "Appointment Facility: Example Clinic\n03/20/2026",
        "progress": "Progress Note: AVERY EXAMPLE, MD 03/20/2026",
    }
    page = "\n".join(fields.values()) + "\nGenerated for Printing on: 03/23/2026"
    quote = fields[quoted_role]
    assert supports_encounter_date(quote, "03/20/2026", page)
    assert not supports_encounter_date(quote, "04/14/1980", page)
    assert not supports_encounter_date(quote, "03/23/2026", page)
    assert not supports_encounter_date(
        quote, "03/20/2026", page.replace("DOS:", "History:")
    )
    assert not supports_encounter_date(
        quote, "03/20/2026", page.replace("MD 03/20/2026", "MD 03/21/2026")
    )
    assert not supports_encounter_date(
        fields["progress"], "03/20/2026", fields["progress"]
    )


@pytest.mark.parametrize(
    "source,quote",
    [
        (
            "Levels L2-L3, L3-L4, L4-L5, and L5-S1.",
            "Levels L2-L3, L3-L4, L4-L5 and L5-S1.",
        ),
        ("Pain and weakness.", "Pain, and weakness."),
        ("Self-\nperformed testing negative.", "Self-performed testing negative."),
        ("NECK: limited range\nof motion.", "Neck: limited range of motion."),
    ],
)
def test_citations_preserve_actual_source_span_after_presentation_matching(
    source, quote
):
    from src.medical_evidence import exact_evidence

    result = exact_evidence(
        [{"page": 2, "quote": quote}], {2: "Header\n" + source + "\nFooter"}
    )
    assert result == [{"page": 2, "quote": source}]


@pytest.mark.parametrize(
    "source,quote",
    [
        ("Dose 1,000 mg.", "Dose 1000 mg."),
        ("Dose 1.0 mg.", "Dose 10 mg."),
        ("No weakness.", "Weakness present."),
        ("Left only.", "Bilateral."),
        ("Level L4-L5.", "Level L4 L5."),
        ("Pain. Unrelated statement. Weakness.", "Pain. Weakness."),
        ("Testing negative.", "Testing positive."),
    ],
)
def test_citation_presentation_matching_does_not_change_clinical_content(source, quote):
    from src.medical_evidence import exact_evidence

    with pytest.raises(EvidenceReview):
        exact_evidence([{"page": 1, "quote": quote}], {1: source, 2: quote})


@pytest.mark.parametrize("separator", [" – ", " — ", " - "])
def test_composite_service_title_requires_both_exact_parts_on_the_same_page(separator):
    from src.medical_evidence import service_title_supported

    description = "follow-up orthopedic evaluation via telemedicine"
    title = "Progress Note" + separator + description
    page = (
        "Seen today for a "
        + description
        + ".\nProgress Note: Avery Example, MD 03/20/2026"
    )
    assert service_title_supported(title, [page])
    assert not service_title_supported(
        title, [page.split("\n")[0], page.split("\n")[1]]
    )
    assert not service_title_supported(title.replace("telemedicine", "surgery"), [page])
    assert not service_title_supported(
        title.replace("Progress Note", "Procedure Note"), [page]
    )
    assert not service_title_supported(
        "Progress Note – orthopedic – telemedicine", [page]
    )


def test_composite_title_reaches_the_independent_clinical_audit(tmp_path):
    text = (
        TEXT
        + "\nSeen for a follow-up orthopedic evaluation.\nProgress Note: Avery Example, MD 02/05/2026"
    )
    data = response()
    data["sections"][0]["entries"][0][
        "service_name"
    ] = "Progress Note – follow-up orthopedic evaluation"
    prompts = []

    def call(prompt, **kwargs):
        prompts.append(prompt)
        if len(prompts) == 1:
            return json.dumps(data)
        candidate = json.loads(prompt[prompt.index('{"policy"') :])["candidate"]
        return json.dumps(
            {
                "entries": [
                    {
                        "id": e["id"],
                        "verdict": "uncertain",
                        "reason": "The encounter attribution needs source review.",
                    }
                    for e in candidate["entries"]
                ],
                "missing_encounters": [],
            }
        )

    result = extract_group(
        [{"page": 1, "text": text}],
        CASE,
        MEDICAL_POLICY,
        DOC,
        call,
        tmp_path / "extract.json",
    )
    assert len(prompts) == 2
    assert not result["entries"] and result["reviews"]


def test_case_list_uses_consolidated_export_count(store):
    case = create_case(store)
    for identifier in ("source-entry-one", "source-entry-two"):
        store.put(case["id"], "entries", identifier, {"id": identifier})
    assert store.overview(case["id"])["entries_count"] == 2
    store.put(
        case["id"],
        "case",
        "export",
        {"version": "example", "entry_ids": ["merged-entry"]},
    )
    state = store.sessions.load(case["id"])
    store.sessions.mark_phase(state, "header", "complete")
    assert store.list_cases()[0]["entries_count"] == 1
    store.sessions.mark_phase(state, "header", "in_progress")
    assert store.overview(case["id"])["entries_count"] == 2


def test_rerun_does_not_claim_previous_upload_completed_current_job(store):
    case, pipeline, calls = prepared_pipeline(store)
    MedicalRun(store, pipeline, case["id"]).run()
    original_artifacts = store.all(case["id"], "artifacts")
    original_calls = dict(calls)

    def pause_after_reset():
        phases = store.overview(case["id"])["phases"]
        for name in ("header", "summary", "upload"):
            assert phases[name]["status"] == "pending"
            assert not phases[name].get("completed_at")
        return True

    from src.session_state import PauseRequested

    with pytest.raises(PauseRequested):
        MedicalRun(store, pipeline, case["id"], stopping=pause_after_reset).run()
    assert store.all(case["id"], "artifacts") == original_artifacts
    assert calls == original_calls


def test_patient_columns_do_not_invent_an_identity_mismatch():
    from src.medical_evidence import explicit_patient_names

    text = (
        TEXT
        + "\nPatient:\nAddress: Example Street\nPatient: Alex Jordan Example | Policy# TEST-ONLY | Service Date: 02/05/2026\n"
    )
    assert explicit_patient_names(text) == ["Alex Example", "Alex Jordan Example"]
    result = validate_document(
        response(), [{"page": 1, "text": text}], CASE, MEDICAL_POLICY, DOC
    )
    assert result["entries"] and not result["reviews"]
    wrong = text + "Patient: Robin Other | Policy# TEST-ONLY\n"
    result = validate_document(
        response(), [{"page": 1, "text": wrong}], CASE, MEDICAL_POLICY, DOC
    )
    assert not result["entries"]
    assert any(i["kind"] == "patient_identity" for i in result["reviews"])


def test_exclude_document_persists_without_ocr_and_is_reversible(store):
    from src.medical_run import issue_for

    case, pipeline, calls = prepared_pipeline(store)
    cid = case["id"]
    MedicalRun(store, pipeline, cid).run()
    doc = store.all(cid, "documents")[0]
    original = (store.directory(cid) / "input" / "record.pdf").read_bytes()
    old_artifacts = store.all(cid, "artifacts")
    for kind in ("coverage", "date"):
        issue = issue_for(doc, kind, "Fictional unresolved question", [1])
        store.put(cid, "issues", issue["id"], issue)
    store.put(
        cid,
        "issues",
        "unrelated",
        {"id": "unrelated", "document_id": "other", "status": "open", "kind": "date"},
    )
    event = review_source(
        store,
        cid,
        {
            "target": issue["id"],
            "fingerprint": issue["fingerprint"],
            "action": "exclude",
            "reviewer": "Test Reviewer",
        },
    )
    assert not store.all(cid, "entries")
    assert store.all(cid, "issues") == [
        {"id": "unrelated", "document_id": "other", "status": "open", "kind": "date"}
    ]
    assert store.all(cid, "artifacts") == old_artifacts
    assert store.get(cid, "document_history", doc["id"] + "-" + event["id"])["entries"]
    # Remove only the synthetic unrelated issue before exporting the fixture.
    with store.connect() as db:
        db.execute(
            "DELETE FROM records WHERE case_id=? AND kind='issues' AND id='unrelated'",
            (cid,),
        )
    before = dict(calls)
    MedicalRun(store, pipeline, cid).run()
    assert calls["ocr"] == before["ocr"] and calls["model"] == before["model"]
    excluded = store.all(cid, "documents")[0]
    assert excluded["status"] == "excluded"
    output = json.loads(
        (store.directory(cid) / "output" / "chronology.json").read_text()
    )
    assert not output["records"] and not output["manual_reviews"]
    assert output["excluded_materials"][0]["reviewer"] == "Test Reviewer"
    review_source(
        store,
        cid,
        {
            "target": excluded["id"],
            "fingerprint": excluded["fingerprint"],
            "action": "rerun_include",
            "reviewer": "Test Reviewer",
        },
    )
    MedicalRun(store, pipeline, cid).run()
    assert store.all(cid, "entries")
    assert calls["ocr"] == before["ocr"] and calls["model"] > before["model"]
    assert (store.directory(cid) / "input" / "record.pdf").read_bytes() == original
    assert {h["action"] for h in store.history(cid)} == {"exclude", "rerun_include"}


def test_exclusion_cannot_hide_a_changed_original(store):
    case, pipeline, _ = prepared_pipeline(store)
    cid = case["id"]
    MedicalRun(store, pipeline, cid).run()
    doc = store.all(cid, "documents")[0]
    review_source(
        store,
        cid,
        {
            "target": doc["id"],
            "fingerprint": doc["sha256"],
            "action": "exclude",
            "reviewer": "Test Reviewer",
        },
    )
    (store.directory(cid) / "input" / "record.pdf").write_bytes(b"changed")
    with pytest.raises(ValueError, match="different source version"):
        MedicalRun(store, pipeline, cid).generate_document(doc)
