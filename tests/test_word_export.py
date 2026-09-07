import io
from docx import Document
from src.word_export import chronology_docx
from src.word_export import output_zip
from zipfile import ZipFile


def test_export_preserves_entries_and_separates_only_explicit_billing():
    text = ('MEDICAL RECORDS SUMMARY\nExample\n\n'
            '03/29/2024–04/17/2024. Clinic. Therapy. C4-C5 unchanged; dose 5–10 mg.\n\n'
            '09/17/2024. Clinic. Procedure (billing record only). Billed service.\n\n'
            '09/17/2024. Clinic. Procedure note. Actual clinical findings.')
    doc = Document(io.BytesIO(chronology_docx(text, separate_billing=True)))
    result = '\n'.join(p.text for p in doc.paragraphs)
    assert '03/29/2024 to 04/17/2024' in result
    assert 'C4-C5 unchanged; dose 5–10 mg.' in result
    assert result.count('Billed service.') == 1
    assert result.index('Actual clinical findings.') < result.index('Billing-only appendix')
    assert result.index('Billed service.') > result.index('Billing-only appendix')
    original = Document(io.BytesIO(chronology_docx(text)))
    assert 'Billing-only appendix' not in '\n'.join(p.text for p in original.paragraphs)


def test_deposition_paragraph_remains_in_dated_chronology():
    entry = '07/16/2026. Jamie Example, Deposition. The patient recalled temporary relief and denied further trauma.'
    doc = Document(io.BytesIO(chronology_docx(entry, separate_billing=True)))
    assert entry in [p.text for p in doc.paragraphs]


def test_word_format_and_service_label():
    doc = Document(io.BytesIO(chronology_docx('05/22/2025. Clinic. Alex Example, MD. Visit Type: Lumbar Block. History: Back pain.')))
    assert doc.styles['Normal'].font.name == 'Times New Roman'
    assert doc.styles['Normal'].font.size.pt == 12
    assert 'Visit Type:' not in doc.paragraphs[0].text
    assert 'Lumbar Block.' in doc.paragraphs[0].text
    assert doc.paragraphs[0].paragraph_format.first_line_indent.inches == .5
    assert 'PAGE' in doc.sections[0].footer._element.xml


def test_zip_handles_binary_word_and_old_sessions(tmp_path):
    markdown = '05/22/2025. Clinic. Office Visit. History: Back pain.'
    (tmp_path/'chronology.md').write_text(markdown)
    for existing in [False, True]:
        if existing:
            (tmp_path/'chronology.docx').write_bytes(chronology_docx(markdown))
        with ZipFile(io.BytesIO(output_zip(tmp_path))) as archive:
            assert archive.namelist().count('chronology.docx') == 1
            doc = Document(io.BytesIO(archive.read('chronology.docx')))
            assert doc.paragraphs[0].text == markdown
            assert archive.read('chronology.md').decode() == markdown
