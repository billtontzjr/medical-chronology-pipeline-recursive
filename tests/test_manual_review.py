from pathlib import Path
"""Source uncertainty produces an explicitly incomplete, reviewable draft."""
import io
import json
import logging
from types import SimpleNamespace

import pytest
from docx import Document

from src.chronology_agent import ChronologyAgent
from src.chronology_scope import ScopeReviewRequired
from src.manual_review import parse_draft_response, collect_manual_reviews, review_markdown, DRAFT_LABEL
from src.scope_recovery import screen_batch
from src.pipeline import MedicalChronologyPipeline
from src.session_state import SessionStore, PHASE_HEADER

DOCS = [{'filename': 'chiro.txt (part 2)', 'source_file': 'chiro.txt', 'content': 'unclear date'},
        {'filename': 'imaging.txt', 'content': 'Date of service: 06/04/2026\nFacility: Imaging\n'
         'Provider: Radiologist, MD\nStudy: MRI\nImpression: Disc protrusion.'}]
CLEAR = '06/04/2026. Imaging. Radiologist, MD. MRI. Impression: Disc protrusion.'
UNCLEAR = '07/12/2014. Clinic. Provider unclear. Office Visit. Back pain.'
DATA = {'sources': [{'id': 'D001', 'scope': 'review_required', 'reason': 'Page 48 date unclear.'},
                    {'id': 'D002', 'scope': 'medical'}],
        'entries': [{'record_type': 'clinical_care', 'source_ids': ['D001'], 'text': UNCLEAR},
                    {'record_type': 'diagnostic_test', 'source_ids': ['D002'],
                     'diagnostic_result': {'date': '06/04/2026', 'facility': 'Imaging',
                         'provider': 'Radiologist, MD', 'study': 'MRI',
                         'evidence': [{'source_id': 'D002', 'date_quote': 'Date of service: 06/04/2026',
                                       'quote': 'Impression: Disc protrusion.'}]}}]}


def agent():
    a = ChronologyAgent.__new__(ChronologyAgent)
    a.logger = logging.getLogger('synthetic'); a.model = 'gpt-6-astra'
    return a


def test_uncertain_source_is_separate_from_medical_draft_and_exclusions():
    text, exclusions, reviews = parse_draft_response(json.dumps(DATA), DOCS)
    assert text == CLEAR and exclusions == []
    assert len(reviews) == 1 and reviews[0]['status'] == 'pending'
    assert reviews[0]['reported_pages'] == [48]
    assert reviews[0]['withheld_proposed_entries'][0]['text'] == UNCLEAR
    assert 'reviewer' in reviews[0] and 'review_date' in reviews[0]


def test_blocked_legacy_checkpoint_recovers_without_api_or_losing_rejection(tmp_path):
    path = tmp_path / 'batch.scope-work.json'
    with pytest.raises(ScopeReviewRequired):
        screen_batch('same prompt', DOCS, lambda *a, **k: json.dumps(DATA), checkpoint=path)
    before = json.loads(path.read_text())
    result = screen_batch('same prompt', DOCS, lambda *a, **k: pytest.fail('reuse saved response'),
                          checkpoint=path, allow_manual_review=True)
    after = json.loads(path.read_text())
    assert result == (CLEAR, [])
    assert after['attempts'] == before['attempts']
    assert after['signature'] == before['signature']
    assert after['status'] == 'complete_with_review'
    assert 'not approved' in after['resolution']
    assert screen_batch('same prompt', DOCS, lambda *a, **k: pytest.fail('still reuse'),
                        checkpoint=path, allow_manual_review=True) == result


@pytest.mark.parametrize('mutation', ['unknown_id', 'duplicate_id', 'missing_source', 'no_reason', 'unknown_entry_id'])
def test_review_flag_cannot_bypass_source_integrity(mutation):
    data = json.loads(json.dumps(DATA))
    if mutation == 'unknown_id': data['sources'][0]['id'] = 'D099'
    if mutation == 'duplicate_id': data['sources'].append(data['sources'][0])
    if mutation == 'missing_source': data['sources'].pop()
    if mutation == 'no_reason': data['sources'][0]['reason'] = ''
    if mutation == 'unknown_entry_id': data['entries'][0]['source_ids'] = ['D099']
    with pytest.raises(ScopeReviewRequired): parse_draft_response(json.dumps(data), DOCS)


def test_flags_are_saved_with_batch_and_resume_without_model(tmp_path):
    a = agent(); a._read_extracted_files = lambda _: DOCS
    a._call_api_with_retry = lambda *a, **k: json.dumps(DATA)
    batches = tmp_path / 'batches'
    a.generate_batches('', str(batches))
    assert (batches / 'batch_001.md').read_text() == CLEAR
    assert len(collect_manual_reviews(batches)) == 1
    a._call_api_with_retry = lambda *a, **k: pytest.fail('completed draft must skip')
    assert a.generate_batches('', str(batches))['batches_skipped_from_disk'] == 1


def assembled(tmp_path, all_uncertain=False):
    a = agent(); batches = tmp_path / 'batches'; output = tmp_path / 'output'; sources = tmp_path / 'sources'
    sources.mkdir(); batches.mkdir()
    text, _, reviews = parse_draft_response(json.dumps(DATA), DOCS)
    (batches / 'batch_001.md').write_text('' if all_uncertain else text)
    (batches / 'batch_001.scope.json').write_text(json.dumps({'exclusions': [], 'manual_reviews': reviews}))
    a._call_api_with_retry = lambda *a, **k: pytest.fail('no merge call needed')
    a._read_extracted_files = lambda _: DOCS
    a.extract_header = lambda *a: {'patient_name': 'SYNTHETIC', 'date_of_birth': '[See Records]', 'date_of_injury': '[See Records]'}
    a.generate_summary_and_gaps = lambda *a: {'summary_md': 'Imaging summary.', 'gaps_md': 'Gaps.'}
    result = a.assemble_outputs(str(sources), str(batches), str(output))
    return result, output


def test_word_markdown_json_and_review_exports_agree(tmp_path):
    result, output = assembled(tmp_path)
    assert result['manual_review_count'] == 1
    md = (output / 'chronology.md').read_text()
    assert DRAFT_LABEL in md and CLEAR in md and UNCLEAR not in md
    word_text = '\n'.join(p.text for p in Document(str(output / 'chronology.docx')).paragraphs)
    assert DRAFT_LABEL in word_text and CLEAR in word_text and UNCLEAR not in word_text
    review_word = '\n'.join(p.text for p in Document(str(output / 'manual_review.docx')).paragraphs)
    assert 'Page 48' in review_word and 'NOT VERIFIED' in review_word and UNCLEAR in review_word
    data = json.loads((output / 'chronology.json').read_text())
    assert data['chronology_markdown'] == md and data['manual_review_required']
    assert not data['excluded_materials']
    assert DRAFT_LABEL in (output / 'summary.md').read_text()
    assert 'Page 48' in (output / 'gaps.md').read_text()


def test_all_deferred_material_still_produces_explicitly_incomplete_draft(tmp_path):
    result, output = assembled(tmp_path, all_uncertain=True)
    assert result['success'] and result['manual_review_count'] == 1
    assert 'No dated entries' in (output / 'summary.md').read_text()
    assert 'nonmedical' not in (output / 'gaps.md').read_text()


def test_verification_retains_pending_review_even_when_model_says_no_issues(tmp_path):
    p = MedicalChronologyPipeline.__new__(MedicalChronologyPipeline)
    p.store = SessionStore(str(tmp_path))
    state = p.store.create(session_id='synthetic', patient_id='synthetic', dropbox_link='', destination_folder='/test')
    _, _, reviews = parse_draft_response(json.dumps(DATA), DOCS)
    (p.store.batches_dir('synthetic') / 'batch_001.scope.json').write_text(json.dumps({'manual_reviews': reviews}))
    output = p.store.output_dir('synthetic'); (output / 'chronology.md').write_text(CLEAR)
    p.chronology_agent = SimpleNamespace(verify_chronology=lambda **kw: {'success': True, 'verification': 'No issues found.'})
    result = p.verify_session('synthetic')
    assert result['manual_review_count'] == 1
    report = (output / 'verification.md').read_text()
    assert DRAFT_LABEL in report and 'Page 48' in report and 'No issues found.' in report
    assert collect_manual_reviews(p.store.batches_dir('synthetic'))[0]['status'] == 'pending'
    (output / 'chronology.md').write_text(DRAFT_LABEL)
    p.chronology_agent = SimpleNamespace(verify_chronology=lambda **kw: pytest.fail('no dated entries to verify'))
    assert p.verify_session('synthetic')['manual_review_count'] == 1
    assert 'No dated entries' in (output / 'verification.md').read_text()


def test_completed_session_ui_shows_manual_review_status(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    import streamlit as st
    from src.session_state import STATUS_COMPLETE
    store = SessionStore(str(tmp_path))
    state = store.create(session_id='synthetic', patient_id='synthetic', dropbox_link='', destination_folder='/test')
    for phase in state.phases.values(): phase.status = STATUS_COMPLETE
    state.status = STATUS_COMPLETE
    state.phases[PHASE_HEADER].data['manual_review_count'] = 1
    store.save(state)
    (store.output_dir('synthetic') / 'chronology.md').write_text(DRAFT_LABEL + '\n\n' + CLEAR)
    (store.output_dir('synthetic') / 'manual_review.md').write_text('Pending page 48 review.')
    def init(self, **kwargs):
        self.store = store
        self.logger = logging.getLogger('synthetic')
        self.chronology_agent = agent()
    monkeypatch.setattr(MedicalChronologyPipeline, '__init__', init)
    for key in ('DROPBOX_APP_KEY','DROPBOX_APP_SECRET','DROPBOX_REFRESH_TOKEN','GOOGLE_CLOUD_API_KEY','ANTHROPIC_API_KEY'):
        monkeypatch.setenv(key, 'synthetic-not-real')
    monkeypatch.setenv('TEAM_PASSWORD', 'synthetic-team-password')
    st.cache_resource.clear()
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py')); app.query_params['session_id'] = 'synthetic'
    app.run()
    app.text_input[0].set_value('synthetic-team-password')
    app.button[0].click().run()
    assert not app.exception
    assert any(DRAFT_LABEL in x.value for x in app.warning)
    assert any(DRAFT_LABEL in x.value for x in app.caption)
    assert any('manual_review.md' in x.proto.label for x in app.get('download_button'))
    st.cache_resource.clear()

