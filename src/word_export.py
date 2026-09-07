"""Deterministic Word export of an existing chronology; no AI calls or factual edits."""
import io
import re
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .encounters import clean_labels

DATE = r'\d{1,2}/\d{1,2}/\d{4}'
BILLING_LABEL = re.compile(r'\(billing record only\)', re.I)


def _display_text(text: str) -> str:
    # Change only full date-to-date separators, never spinal levels or doses.
    text = clean_labels(text)
    return re.sub(rf'({DATE})\s*[–—−-]\s*({DATE})', r'\1 to \2', text)


def chronology_docx(markdown: str, *, separate_billing: bool = False) -> bytes:
    """Preserve every entry, optionally relocate explicitly labeled billing entries.

    The label is taken from the existing draft, not independently validated.
    Paragraph conversion deliberately leaves unfamiliar syntax/text intact.
    """
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.5)
    # Standard business brief body, with a restrained clinical heading override.
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(10)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.widow_control = True
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
    doc.core_properties.author = ''
    doc.core_properties.title = 'Medical chronology'
    doc.core_properties.comments = ''
    blocks = [b.strip() for b in re.split(r'\n\s*\n', markdown.strip()) if b.strip()]
    clinical, billing = [], []
    for block in blocks:
        # Only explicit labels in a dated entry header qualify for relocation.
        header = block.splitlines()[0]
        target = billing if (separate_billing and re.match(DATE, header)
                             and BILLING_LABEL.search(header)) else clinical
        target.append(block)

    def add_block(block):
        if block.startswith('MEDICAL RECORDS SUMMARY'):
            doc.add_paragraph()
            for line in block.splitlines():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0 if line.startswith('Date of Birth:') else 10)
                p.add_run(line).bold = True
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(.5)
        for i, line in enumerate(block.splitlines()):
            if i:
                p.add_run().add_break()
            # Preserve inline emphasis without emitting Markdown stars.
            line = re.sub(r'^#{1,6}\s+', '', line)
            for piece in re.split(r'(\*\*[^*]+\*\*)', _display_text(line)):
                if piece.startswith('**') and piece.endswith('**'):
                    p.add_run(piece[2:-2]).bold = True
                else:
                    p.add_run(piece)
    for block in clinical:
        add_block(block)
    if billing:
        doc.add_page_break()
        doc.add_paragraph('Billing-only appendix', style='Heading 1')
        doc.add_paragraph('These entries were labeled "billing record only" in the generated draft. '
                          'Their presence does not establish that clinical notes are missing. '
                          'Use them to reconcile bills with the source records.')
        for block in billing:
            add_block(block)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def output_zip(output_dir):
    """Package text and binary artifacts; supply Word for older saved sessions."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(output_dir.iterdir()):
            if path.is_file():
                archive.writestr(path.name, path.read_bytes())
        markdown = output_dir / 'chronology.md'
        if markdown.exists() and not (output_dir / 'chronology.docx').exists():
            archive.writestr('chronology.docx', chronology_docx(markdown.read_text(encoding='utf-8')))
    return output.getvalue()
